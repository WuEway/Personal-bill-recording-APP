from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

import openpyxl

from mz.importer.base import BaseImporter
from mz.models.transaction import RawTransactionDraft

HEADER_MARKER = "----------------------微信支付账单明细列表--------------------"
EXPECTED_HEADER = (
    "交易时间", "交易类型", "交易对方", "商品", "收/支",
    "金额(元)", "支付方式", "当前状态", "交易单号", "商户单号", "备注",
)


def _parse_amount(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().lstrip("¥").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


class WechatImporter(BaseImporter):
    SOURCE_NAME = "wechat"
    DISPLAY_NAME = "微信支付账单"
    SUPPORTED_FORMATS = ["xlsx", "csv"]

    def detect(self, file_path: Path, file_format: str) -> bool:
        if file_format == "xlsx":
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws = wb.active
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 20:
                        break
                    if row and row[0] and "微信支付账单明细" in str(row[0]):
                        return True
                wb.close()
            except Exception:
                pass
        elif file_format == "csv":
            try:
                text = file_path.read_text(encoding="utf-8", errors="ignore")
                return "微信支付账单明细" in text[:2000]
            except Exception:
                pass
        return False

    def parse(self, file_path: Path, file_format: str) -> Iterator[RawTransactionDraft]:
        if file_format == "xlsx":
            yield from self._parse_xlsx(file_path)
        else:
            yield from self._parse_csv(file_path)

    def detect_period(self, file_path: Path, file_format: str) -> tuple[date, date]:
        dates: list[date] = []
        for draft in self.parse(file_path, file_format):
            dates.append(draft.txn_time.date())
        if not dates:
            today = date.today()
            return today, today
        return min(dates), max(dates)

    def _parse_xlsx(self, file_path: Path) -> Iterator[RawTransactionDraft]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_found = False
        for row in rows:
            if not header_found:
                if row and row[0] and str(row[0]).startswith("交易时间"):
                    header_found = True
                continue
            if not row or not row[0]:
                continue
            draft = self._parse_row(row)
            if draft:
                yield draft
        wb.close()

    def _parse_csv(self, file_path: Path) -> Iterator[RawTransactionDraft]:
        text = file_path.read_text(encoding="utf-8-sig", errors="replace")
        lines = text.splitlines()
        header_idx = None
        for i, line in enumerate(lines):
            if line.startswith("交易时间,"):
                header_idx = i
                break
        if header_idx is None:
            return
        reader = csv.reader(lines[header_idx + 1:])
        for row in reader:
            if not row or not row[0].strip():
                continue
            draft = self._parse_row(tuple(row))
            if draft:
                yield draft

    def _parse_row(self, row: tuple) -> RawTransactionDraft | None:
        try:
            (time_val, txn_type, party, item, sign, amount_val,
             payment, status, txn_id, merchant_id, note) = row[:11]

            # Parse time
            if isinstance(time_val, datetime):
                txn_time = time_val
            else:
                time_str = str(time_val).strip()
                for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y-%m-%d"):
                    try:
                        txn_time = datetime.strptime(time_str, fmt)
                        break
                    except ValueError:
                        continue
                else:
                    return None  # unparseable date

            # Skip header rows that might be repeated in data
            if str(time_val).strip() in ("交易时间", ""):
                return None

            amount = _parse_amount(amount_val)
            sign_str = str(sign).strip() if sign else ""

            if sign_str == "支出":
                direction = "expense"
                amount_cents = -int(round(abs(amount) * 100))
            elif sign_str == "收入":
                direction = "income"
                amount_cents = int(round(abs(amount) * 100))
            else:
                direction = "transfer"
                amount_cents = int(round(amount * 100))

            is_group_payment = str(txn_type).strip() == "群收款"

            raw_dict: dict[str, Any] = {
                k: str(v) if v is not None else None
                for k, v in zip(EXPECTED_HEADER, row[:11])
            }

            return RawTransactionDraft(
                source="wechat",
                txn_time=txn_time,
                amount_cents=amount_cents,
                counterparty=str(party).strip() if party else None,
                description=str(item).strip() if item else None,
                payment_method_raw=str(payment).strip() if payment else None,
                txn_type_raw=str(txn_type).strip() if txn_type else None,
                direction=direction,
                external_txn_id=str(txn_id).strip() if txn_id else None,
                external_merchant_id=str(merchant_id).strip() if merchant_id else None,
                is_group_payment=is_group_payment,
                raw_json=raw_dict,
            )
        except Exception:
            return None
